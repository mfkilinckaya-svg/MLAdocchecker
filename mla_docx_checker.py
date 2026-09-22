"""
MLA DOCX Checker
================
MLA Handbook 9. baskı (2021) biçim kurallarına göre Word (.docx) belgelerini
denetler. Her kontrolün yanında ilgili Handbook bölüm numarası belirtilmiştir.

Kullanım:
    python mla_docx_checker.py makale.docx
"""

import re
import sys
import unicodedata

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Length

__version__ = "2.1.0"

TOL = 0.06  # inç cinsinden tolerans (Word değerleri yuvarlanabiliyor)
WC_HEADINGS = ("Works Cited", "Work Cited")
ARTICLES = re.compile(r"^(a|an|the)\s+", re.IGNORECASE)
TRIPLE_DASH = re.compile(r"^(-{3}|—{3}|–{3})\s*\.?\s*")


def _close(value, target, tol=TOL):
    return value is not None and abs(value - target) <= tol


def _alpha_key(text):
    """Harf harf (letter-by-letter) sıralama anahtarı — MLA 5.124.
    Aksan ve diğer işaretler yok sayılır, baştaki A/An/The atlanır,
    ters çevrilmiş yazar adında önce virgüle kadar olan kısım karşılaştırılır."""
    t = text.replace("ı", "i").replace("İ", "I")
    t = unicodedata.normalize("NFKD", t)
    t = "".join(c for c in t if not unicodedata.combining(c))
    t = t.strip("\"'“”‘’ ").lower()
    t = ARTICLES.sub("", t)
    head, _, tail = t.partition(",")
    clean = lambda s: re.sub(r"[^a-z0-9]", "", s)
    return (clean(head), clean(tail))


class MLADocxChecker:
    """MLA 9. baskı standartlarına göre .docx belgelerini analiz eder."""

    def __init__(self, doc_path):
        self.doc_path = doc_path
        self.doc = docx.Document(doc_path)
        self.report = []
        self.paragraphs = self.doc.paragraphs
        self.wc_index = self._find_works_cited()
        self._defaults = self._read_doc_defaults()

    # ------------------------------------------------------------------ #
    # Yardımcılar
    # ------------------------------------------------------------------ #
    def log(self, category, status, message):
        self.report.append({"kategori": category, "durum": status, "mesaj": message})

    def _find_works_cited(self):
        for i, p in enumerate(self.paragraphs):
            if p.text.strip().rstrip(".") in WC_HEADINGS:
                return i
        return None

    def _read_doc_defaults(self):
        """styles.xml içindeki docDefaults değerlerini okur (stil zinciri boşsa kullanılır)."""
        d = {"line": None, "size": None, "font": None}
        root = self.doc.styles.element
        sp = root.find(f"{qn('w:docDefaults')}/{qn('w:pPrDefault')}/{qn('w:pPr')}/{qn('w:spacing')}")
        if sp is not None and sp.get(qn("w:line")):
            if sp.get(qn("w:lineRule"), "auto") == "auto":
                d["line"] = int(sp.get(qn("w:line"))) / 240
        rpr = root.find(f"{qn('w:docDefaults')}/{qn('w:rPrDefault')}/{qn('w:rPr')}")
        if rpr is not None:
            sz = rpr.find(qn("w:sz"))
            if sz is not None:
                d["size"] = int(sz.get(qn("w:val"))) / 2
            fonts = rpr.find(qn("w:rFonts"))
            if fonts is not None:
                d["font"] = fonts.get(qn("w:ascii"))
        return d

    def _style_chain(self, style):
        while style is not None:
            yield style
            style = style.base_style

    def _pfmt(self, p, attr):
        """Paragraf biçim özelliğini önce paragraftan, sonra stil zincirinden çözer."""
        val = getattr(p.paragraph_format, attr)
        if val is not None:
            return val
        for st in self._style_chain(p.style):
            val = getattr(st.paragraph_format, attr)
            if val is not None:
                return val
        return None

    def _alignment(self, p):
        return self._pfmt(p, "alignment")

    def _line_spacing(self, p):
        """(değer, kural) döndürür. Kat (multiple) aralıkta değer float olur."""
        val = self._pfmt(p, "line_spacing")
        rule = self._pfmt(p, "line_spacing_rule")
        if val is None:
            val = self._defaults["line"] or 1.0
        return val, rule

    def _is_double(self, p):
        val, rule = self._line_spacing(p)
        if isinstance(val, Length) or rule in (WD_LINE_SPACING.EXACTLY, WD_LINE_SPACING.AT_LEAST):
            return False, "sabit"
        if rule == WD_LINE_SPACING.DOUBLE or _close(float(val), 2.0, 0.05):
            return True, None
        return False, round(float(val), 2)

    def _inches(self, p, attr):
        v = self._pfmt(p, attr)
        return round(v.inches, 2) if v is not None else 0.0

    def _run_font(self, run, p):
        if run.font.name:
            return run.font.name
        for st in self._style_chain(p.style):
            if st.font.name:
                return st.font.name
        return self._defaults["font"]

    def _run_size(self, run, p):
        if run.font.size:
            return run.font.size.pt
        for st in self._style_chain(p.style):
            if st.font.size:
                return st.font.size.pt
        return self._defaults["size"]

    def _is_heading(self, p):
        return p.style is not None and p.style.name.lower().startswith(("heading", "başlık"))

    def _has_page_break_before(self, idx):
        p = self.paragraphs[idx]
        if self._pfmt(p, "page_break_before") or "lastRenderedPageBreak" in p._p.xml:
            return True
        for j in range(idx - 1, -1, -1):
            prev = self.paragraphs[j]
            xml = prev._p.xml
            if 'w:type="page"' in xml or "<w:sectPr" in xml:
                return True
            if prev.text.strip():
                break
        return False

    def _body_range(self):
        """Başlıktan (5. dolu paragraf) sonra, Works Cited'e kadar olan dizin aralığı."""
        filled = [i for i, p in enumerate(self.paragraphs) if p.text.strip()]
        start = filled[5] if len(filled) > 5 else len(self.paragraphs)
        end = self.wc_index if self.wc_index is not None else len(self.paragraphs)
        return start, end

    # ------------------------------------------------------------------ #
    # Kontroller
    # ------------------------------------------------------------------ #
    def run_all_checks(self):
        self.check_margins()
        self.check_running_head()
        self.check_hyphenation()
        self.check_first_page_and_title()
        self.check_text_formatting()
        self.check_headings()
        self.check_in_text_citations()
        self.check_works_cited()
        return self.report

    # 1.1 — Kenar boşlukları
    def check_margins(self):
        for idx, s in enumerate(self.doc.sections, 1):
            m = {k: round(getattr(s, f"{k}_margin").inches, 2) for k in ("top", "bottom", "left", "right")}
            if all(_close(v, 1.0) for v in m.values()):
                self.log("Kenar Boşlukları", "BAŞARILI", f"Bölüm {idx}: Tüm kenar boşlukları 1 inç (2,54 cm).")
            else:
                self.log("Kenar Boşlukları", "HATA",
                         f"Bölüm {idx}: Kenar boşlukları 1 inç değil (Üst={m['top']}\", Alt={m['bottom']}\", "
                         f"Sol={m['left']}\", Sağ={m['right']}\"). [MLA 1.1]")

    # 1.4 — Üst bilgi: soyadı + sayfa numarası, sağ üstte, üstten 0,5 inç
    def check_running_head(self):
        for idx, s in enumerate(self.doc.sections, 1):
            header = s.header
            if idx > 1 and header.is_linked_to_previous:
                continue
            xml = header._element.xml
            text = " ".join(p.text for p in header.paragraphs).strip()
            if "PAGE" not in xml:
                self.log("Üst Bilgi (Running Head)", "HATA",
                         f"Bölüm {idx}: Üst bilgide otomatik sayfa numarası yok. Her sayfanın sağ üstünde "
                         f"'Soyadı SayfaNo' bulunmalıdır. [MLA 1.4]")
                continue
            ok = True
            if not text:
                ok = False
                self.log("Üst Bilgi (Running Head)", "UYARI",
                         f"Bölüm {idx}: Sayfa numarasının önünde soyadı bulunamadı. [MLA 1.4]")
            if re.search(r"\b(p|pg|page|sayfa)\.?\s*$", text, re.IGNORECASE) or re.search(r"[-–.]\s*$", text):
                ok = False
                self.log("Üst Bilgi (Running Head)", "HATA",
                         f"Bölüm {idx}: Sayfa numarasının önüne 'p.', tire, nokta vb. konmaz. [MLA 1.4]")
            aligns = {self._alignment(p) for p in header.paragraphs if p.text.strip() or "PAGE" in p._p.xml}
            if WD_ALIGN_PARAGRAPH.RIGHT not in aligns:
                ok = False
                self.log("Üst Bilgi (Running Head)", "HATA", f"Bölüm {idx}: Üst bilgi sağa hizalı değil. [MLA 1.4]")
            hd = s.header_distance.inches if s.header_distance is not None else None
            if not _close(hd, 0.5):
                ok = False
                self.log("Üst Bilgi (Running Head)", "UYARI",
                         f"Bölüm {idx}: Üst bilgi sayfanın üstünden 0,5 inç uzakta olmalıdır. [MLA 1.4]")
            if s.different_first_page_header_footer and "PAGE" not in s.first_page_header._element.xml:
                ok = False
                self.log("Üst Bilgi (Running Head)", "HATA",
                         f"Bölüm {idx}: 'Farklı ilk sayfa' açık ve ilk sayfada numara yok; üst bilgi ilk sayfa "
                         f"dahil her sayfada olmalıdır. [MLA 1.4]")
            if ok:
                self.log("Üst Bilgi (Running Head)", "BAŞARILI", f"Bölüm {idx}: Soyadı + sayfa numarası sağ üstte.")

    # 1.2 — Otomatik heceleme kapalı olmalı
    def check_hyphenation(self):
        el = self.doc.settings.element.find(qn("w:autoHyphenation"))
        if el is not None and el.get(qn("w:val"), "true") not in ("false", "0", "off"):
            self.log("Heceleme", "HATA", "Otomatik heceleme (hyphenation) açık; kapatılmalıdır. [MLA 1.2]")

    # 1.3 — İlk sayfa kimlik bloğu ve başlık
    def check_first_page_and_title(self):
        filled = [p for p in self.paragraphs if p.text.strip()]
        if len(filled) < 5:
            self.log("İlk Sayfa Düzeni", "UYARI", "Belgede yeterli paragraf bulunamadı.")
            return

        for i, p in enumerate(filled[:4], 1):
            if self._alignment(p) not in (None, WD_ALIGN_PARAGRAPH.LEFT):
                self.log("İlk Sayfa Kimlik Bloğu", "HATA",
                         f"Kimlik satırı {i} ('{p.text.strip()}') sola hizalı olmalıdır. [MLA 1.3]")
            if abs(self._inches(p, "first_line_indent")) > TOL or abs(self._inches(p, "left_indent")) > TOL:
                self.log("İlk Sayfa Kimlik Bloğu", "HATA",
                         f"Kimlik satırı {i} ('{p.text.strip()}') girintisiz, sol kenara bitişik olmalıdır. [MLA 1.3]")
        if not re.search(r"\d", filled[3].text):
            self.log("İlk Sayfa Kimlik Bloğu", "UYARI",
                     f"4. satır tarih gibi görünmüyor ('{filled[3].text.strip()}'). Sıra: ad, öğretim üyesi, "
                     f"ders, tarih (ör. 21 October 2019). [MLA 1.3]")

        title_p = filled[4]
        title = title_p.text.strip()
        if self._alignment(title_p) != WD_ALIGN_PARAGRAPH.CENTER:
            self.log("Makale Başlığı", "HATA", f"Başlık ('{title}') ortalanmamış. [MLA 1.3]")
        else:
            self.log("Makale Başlığı", "BAŞARILI", f"Başlık ortalanmış: '{title}'")
        if title.endswith("."):
            self.log("Makale Başlığı", "HATA", f"Başlığın sonuna nokta konmaz: '{title}' [MLA 1.3]")
        if title[:1] in "\"“'‘" and title[-1:] in "\"”'’":
            self.log("Makale Başlığı", "HATA", "Başlık tırnak içine alınmamalıdır. [MLA 1.3]")
        letters = [c for c in title if c.isalpha()]
        if len(letters) > 3 and all(c.isupper() for c in letters):
            self.log("Makale Başlığı", "HATA", "Başlık tamamen büyük harfle yazılmamalıdır. [MLA 1.3]")

        runs = [r for r in title_p.runs if r.text.strip()]
        if any(r.bold for r in runs):
            self.log("Makale Başlığı Biçimi", "HATA", "Başlık kalın (bold) olmamalıdır. [MLA 1.3]")
        if any(r.underline for r in runs):
            self.log("Makale Başlığı Biçimi", "HATA", "Başlığın altı çizilmemelidir. [MLA 1.3]")
        if runs and all(r.italic for r in runs):
            self.log("Makale Başlığı Biçimi", "HATA",
                     "Başlığın tamamı italik olmamalıdır; yalnızca içindeki eser adları italik yazılır. [MLA 1.3]")

    # 1.2 — Yazı tipi, boyut, satır aralığı, hizalama, girinti
    def check_text_formatting(self):
        body_start, body_end = self._body_range()
        limit = self.wc_index if self.wc_index is not None else len(self.paragraphs)

        issues = {k: [] for k in ("justify", "spacing", "exact", "space", "indent")}
        fonts, sizes, bad_sizes = set(), set(), set()
        double_spaces, blank_lines = 0, 0

        for idx, p in enumerate(self.paragraphs[:limit]):
            text = p.text.strip()
            n = idx + 1
            if not text:
                if body_start <= idx < body_end and not p._p.xpath(".//w:br | .//w:drawing"):
                    blank_lines += 1
                continue

            if self._alignment(p) == WD_ALIGN_PARAGRAPH.JUSTIFY:
                issues["justify"].append(n)

            ok, val = self._is_double(p)
            if not ok:
                (issues["exact"] if val == "sabit" else issues["spacing"]).append(
                    n if val == "sabit" else f"{n} ({val})")

            sb, sa = self._pfmt(p, "space_before"), self._pfmt(p, "space_after")
            if (sb and sb.pt > 0) or (sa and sa.pt > 0):
                issues["space"].append(n)

            # Gövdede 0,5" ilk satır girintisi (ara başlıklar ve blok alıntılar hariç)
            if body_start <= idx < body_end and not self._is_heading(p):
                left = self._inches(p, "left_indent")
                first = self._inches(p, "first_line_indent")
                is_block_quote = _close(left, 0.5) and _close(first, 0.0)
                centered = self._alignment(p) in (WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT)
                if not is_block_quote and not centered and not _close(first, 0.5):
                    issues["indent"].append(n)

            if re.search(r"[.?!]\s{2,}\S", p.text):
                double_spaces += 1

            for r in p.runs:
                if not r.text.strip():
                    continue
                f, s = self._run_font(r, p), self._run_size(r, p)
                if f:
                    fonts.add(f)
                if s:
                    sizes.add(s)
                    if not (11 <= s <= 13):
                        bad_sizes.add(s)

        def short(lst, k=12):
            return ", ".join(map(str, lst[:k])) + (" ..." if len(lst) > k else "")

        if issues["justify"]:
            self.log("Metin Hizalama", "HATA",
                     f"İki yana yaslanmış paragraflar: {short(issues['justify'])}. Metin sola hizalanmalıdır. [MLA 1.2]")
        if issues["exact"]:
            self.log("Satır Aralığı", "HATA",
                     f"Sabit (tam/en az) satır aralığı kullanılan paragraflar: {short(issues['exact'])}. "
                     f"'Çift' (2,0) seçilmelidir. [MLA 1.2]")
        if issues["spacing"]:
            self.log("Satır Aralığı", "HATA",
                     f"Çift satır aralıklı olmayan paragraflar: {short(issues['spacing'])}. [MLA 1.2]")
        if not issues["exact"] and not issues["spacing"]:
            self.log("Satır Aralığı", "BAŞARILI", "Tüm metin çift satır aralıklı.")
        if issues["space"]:
            self.log("Paragraf Boşlukları", "UYARI",
                     f"Öncesinde/sonrasında ekstra boşluk olan paragraflar: {short(issues['space'])}. "
                     f"Çift aralık dışında ek boşluk bırakılmamalıdır (0 pt). [MLA 1.2]")
        if issues["indent"]:
            self.log("Paragraf Girintisi", "HATA",
                     f"İlk satırı 0,5 inç girintili olmayan gövde paragrafları: {short(issues['indent'])}. [MLA 1.2]")
        if blank_lines:
            self.log("Boş Satırlar", "UYARI",
                     f"Gövdede {blank_lines} boş paragraf var. Paragraflar arasında boş satır bırakılmaz; "
                     f"aralık yalnızca çift satır aralığıyla sağlanır. [MLA 1.2]")
        if double_spaces:
            self.log("Cümle Arası Boşluk", "UYARI",
                     f"{double_spaces} paragrafta cümle sonundan sonra iki boşluk var; öğretim üyesi aksini "
                     f"istemedikçe tek boşluk bırakılır. [MLA 1.2]")

        if bad_sizes:
            self.log("Yazı Boyutu", "HATA", f"11–13 punto aralığı dışında boyutlar: {sorted(bad_sizes)} pt. [MLA 1.2]")
        if len(sizes) > 1:
            self.log("Yazı Boyutu", "UYARI",
                     f"Birden fazla yazı boyutu kullanılmış: {sorted(sizes)} pt. Belge boyunca tek boyut kullanın. [MLA 1.2]")
        if len(fonts) > 1:
            self.log("Yazı Tipi", "UYARI",
                     f"Birden fazla yazı tipi kullanılmış: {', '.join(sorted(fonts))}. Tek, okunaklı bir yazı tipi "
                     f"(ör. Times New Roman) kullanın. [MLA 1.2]")
        elif fonts:
            self.log("Yazı Tipi", "BAŞARILI", f"Tek yazı tipi kullanılmış: {next(iter(fonts))}.")

    # 1.5 — Ara başlıklar
    def check_headings(self):
        start, end = self._body_range()
        for idx in range(start, end):
            p = self.paragraphs[idx]
            t = p.text.strip()
            if not t or not self._is_heading(p):
                continue
            if t.endswith("."):
                self.log("Ara Başlık", "HATA", f"Başlıktan sonra nokta konmaz: '{t}' [MLA 1.3, 1.5]")
            if self._alignment(p) in (WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT):
                self.log("Ara Başlık", "UYARI", f"Gövdedeki ara başlıklar sola yaslı olmalıdır: '{t}' [MLA 1.5]")
            letters = [c for c in t if c.isalpha()]
            if len(letters) > 3 and all(c.isupper() for c in letters):
                self.log("Ara Başlık", "UYARI", f"Ara başlıkta tamamen büyük harf kullanmayın: '{t}' [MLA 1.5]")

    # 6.x — Metin içi alıntılar
    def check_in_text_citations(self):
        limit = self.wc_index if self.wc_index is not None else len(self.paragraphs)
        name = r"[^\W\d_][^\W\d_'’\-]*"
        author = rf"{name}(?:\s+(?:and\s+)?{name})*(?:\s+et al\.)?"
        pages = r"\d+(?:\s*[-–]\s*\d+)?"
        comma_re = re.compile(rf"\(({author}),\s*({pages})\)")
        period_re = re.compile(rf"\((?:{author}\s+)?{pages}\.\)")
        pp_re = re.compile(rf"\((?:{author},?\s+)?(?:pp?\.|pg\.?)\s*\d[^)]*\)", re.IGNORECASE)
        etal_re = re.compile(r"\bet\.?\s*al\b\.?")

        for idx, p in enumerate(self.paragraphs[:limit], 1):
            text = p.text
            if not text.strip():
                continue
            for a, pg in comma_re.findall(text):
                self.log("Metin İçi Alıntı", "HATA",
                         f"Paragraf {idx}: Yazar ile sayfa numarası arasına virgül konmaz: "
                         f"'({a}, {pg})' → '({a} {pg})' [MLA 6.30]")
            for m in period_re.findall(text):
                self.log("Noktalama Konumu", "HATA",
                         f"Paragraf {idx}: Nokta parantezin dışına konmalıdır: '{m}' [MLA 6.43]")
            for m in pp_re.finditer(text):
                self.log("Metin İçi Alıntı", "HATA",
                         f"Paragraf {idx}: Parantez içi alıntıda sayfa numarasından önce 'p.'/'pp.' kullanılmaz: "
                         f"'{m.group(0)}' [MLA 6.16]")
            for e in etal_re.findall(text):
                if e != "et al.":
                    self.log("Alıntı Biçimi (et al.)", "HATA",
                             f"Paragraf {idx}: 'et al.' hatalı yazılmış ('{e}'); doğru biçim 'et al.' [MLA 5.8, 6.5]")

    # 1.6, 5.120–5.130 — Kaynakça (Works Cited)
    def check_works_cited(self):
        if self.wc_index is None:
            self.log("Kaynakça Sayfası", "UYARI", "'Works Cited' başlığı bulunamadı.")
            return

        head = self.paragraphs[self.wc_index]
        head_text = head.text.strip()
        entries = [p for p in self.paragraphs[self.wc_index + 1:] if p.text.strip()]

        if self._alignment(head) != WD_ALIGN_PARAGRAPH.CENTER:
            self.log("Kaynakça Başlığı", "HATA", "'Works Cited' başlığı ortalanmamış. [MLA 1.6]")
        else:
            self.log("Kaynakça Başlığı", "BAŞARILI", "'Works Cited' başlığı ortalanmış.")
        if head_text.endswith("."):
            self.log("Kaynakça Başlığı", "HATA", "'Works Cited' başlığından sonra nokta konmaz. [MLA 1.3]")
        if any(r.bold or r.italic or r.underline for r in head.runs if r.text.strip()):
            self.log("Kaynakça Başlığı", "UYARI", "'Works Cited' başlığı düz yazılmalıdır (kalın/italik/altı çizili değil).")
        if len(entries) == 1 and head_text.rstrip(".") == "Works Cited":
            self.log("Kaynakça Başlığı", "HATA", "Tek kaynak varsa başlık 'Work Cited' olmalıdır. [MLA 1.6]")
        if len(entries) > 1 and head_text.rstrip(".") == "Work Cited":
            self.log("Kaynakça Başlığı", "HATA", "Birden fazla kaynak varsa başlık 'Works Cited' olmalıdır. [MLA 1.6]")
        if not self._has_page_break_before(self.wc_index):
            self.log("Kaynakça Sayfası", "UYARI",
                     "'Works Cited' yeni bir sayfada başlamıyor gibi görünüyor (sayfa sonu bulunamadı). [MLA 1.6]")

        if not entries:
            self.log("Kaynakça Sayfası", "UYARI", "'Works Cited' altında kaynak bulunamadı.")
            return

        # Alfabetik sıra: harf harf, aksanlar yok sayılır; '---.' önceki yazarı temsil eder,
        # aynı yazarın eserleri başlığa göre sıralanır (5.124–5.126)
        keys, prev_author = [], ("", "")
        for p in entries:
            t = p.text.strip()
            m = TRIPLE_DASH.match(t)
            if m:
                keys.append((prev_author, _alpha_key(t[m.end():])))
            else:
                author_part, _, rest = t.partition(". ")
                prev_author = _alpha_key(author_part)
                keys.append((prev_author, _alpha_key(rest)))
        misplaced = [entries[i].text.strip()[:40] for i in range(1, len(keys)) if keys[i] < keys[i - 1]]
        if misplaced:
            self.log("Kaynakça Sıralaması", "HATA",
                     "Works Cited alfabetik sırada değil. Yerinde olmayan girdi(ler): "
                     + "; ".join(f"'{m}...'" for m in misplaced) + " [MLA 5.124]")
        else:
            self.log("Kaynakça Sıralaması", "BAŞARILI", "Kaynaklar alfabetik olarak sıralı.")

        bad_hang, bad_space, no_period = [], [], []
        for p in entries:
            t = p.text.strip()
            if not (_close(self._inches(p, "left_indent"), 0.5) and _close(self._inches(p, "first_line_indent"), -0.5)):
                bad_hang.append(t[:35])
            if not self._is_double(p)[0]:
                bad_space.append(t[:35])
            if not t.endswith("."):
                no_period.append(t[:35])

        if bad_hang:
            self.log("Asılı Girinti (Hanging Indent)", "HATA",
                     f"{len(bad_hang)} girdide 0,5 inç asılı girinti yok (ör. '{bad_hang[0]}...'). [MLA 1.6]")
        else:
            self.log("Asılı Girinti (Hanging Indent)", "BAŞARILI", "Tüm girdilerde 0,5 inç asılı girinti var.")
        if bad_space:
            self.log("Kaynakça Satır Aralığı", "HATA",
                     f"{len(bad_space)} girdi çift satır aralıklı değil (ör. '{bad_space[0]}...'). [MLA 1.6]")
        if no_period:
            self.log("Kaynakça Noktalama", "UYARI",
                     f"{len(no_period)} girdi nokta ile bitmiyor (ör. '{no_period[0]}...'). "
                     f"Her girdi nokta ile biter. [MLA 5.120]")


def main():
    file_name = sys.argv[1] if len(sys.argv) > 1 else "sample_mla_paper.docx"
    print("=" * 50)
    print(f" MLA 9th Edition Word (.docx) Analiz Raporu  v{__version__}")
    print(f" Belge: {file_name}")
    print("=" * 50 + "\n")

    try:
        report = MLADocxChecker(file_name).run_all_checks()
    except Exception as e:
        print(f"❌ Analiz sırasında bir hata oluştu: {e}")
        sys.exit(2)

    icons = {"BAŞARILI": "✅", "UYARI": "⚠️", "HATA": "❌"}
    stats = {k: 0 for k in icons}
    for item in report:
        stats[item["durum"]] += 1
        print(f"{icons[item['durum']]} [{item['durum']}] {item['kategori']}: {item['mesaj']}")

    print("\n" + "-" * 50)
    print(f"Özet: {stats['BAŞARILI']} Başarılı, {stats['UYARI']} Uyarı, {stats['HATA']} Hata.")
    print("-" * 50)
    sys.exit(1 if stats["HATA"] else 0)


if __name__ == "__main__":
    main()
